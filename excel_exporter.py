import pandas as pd
import logging
from pathlib import Path
from typing import Dict, List, Any
from datetime import datetime
from config import EXCEL_COLUMNS

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class ExcelExporter:
    """Handles Excel export functionality"""
    
    def __init__(self):
        self.columns = EXCEL_COLUMNS
    
    def export_results(self, results: List[Dict[str, Any]], output_path: str = None) -> str:
        """
        Export analysis results to Excel file
        
        Args:
            results: List of dictionaries containing analysis results
            output_path: Optional custom output path
            
        Returns:
            Path to the exported Excel file
        """
        try:
            # Create DataFrame
            df = pd.DataFrame(results)
            
            # Ensure all required columns exist
            for col in self.columns:
                if col not in df.columns:
                    df[col] = ""
            
            # Reorder columns to match specification
            df = df[self.columns]
            
            # Generate output path if not provided
            if output_path is None:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                output_path = f"document_analysis_results_{timestamp}.xlsx"
            
            # Export to Excel with formatting
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name='Document Analysis', index=False)
                
                # Get the workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['Document Analysis']
                
                # Apply formatting
                self._apply_excel_formatting(workbook, worksheet, df)
            
            logger.info(f"Results exported to: {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"Error exporting to Excel: {str(e)}")
            raise
    
    def _apply_excel_formatting(self, workbook, worksheet, df):
        """Apply formatting to the Excel worksheet"""
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        # Header formatting
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Apply header formatting
        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # Column width adjustments
        column_widths = {
            'A': 25,  # file_name
            'B': 15,  # file_type
            'C': 50,  # extracted_text
            'D': 50,  # summary
            'E': 30,  # description
            'F': 30,  # bullet_points
            'G': 12,  # version
            'H': 25,  # tags
            'I': 30,  # duplicates
            'J': 15,  # duplicates_percentage
            'K': 25   # master_one
        }
        
        for col, width in column_widths.items():
            worksheet.column_dimensions[col].width = width
        
        # Apply borders and alignment to all cells
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        cell_alignment = Alignment(vertical="top", wrap_text=True)
        
        for row in worksheet.iter_rows(min_row=1, max_row=len(df) + 1):
            for cell in row:
                cell.border = thin_border
                if cell.row > 1:  # Don't change header alignment
                    cell.alignment = cell_alignment
        
        # Freeze the header row
        worksheet.freeze_panes = "A2"
        
        # Apply conditional formatting for duplicates
        duplicate_fill = PatternFill(start_color="FFE6E6", end_color="FFE6E6", fill_type="solid")
        
        for row_idx, row in enumerate(df.itertuples(), start=2):
            if hasattr(row, 'duplicates_percentage') and row.duplicates_percentage > 0:
                for col_idx in range(1, len(self.columns) + 1):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    cell.fill = duplicate_fill
    
    def prepare_row_data(self, file_info: Dict[str, Any]) -> Dict[str, Any]:
        """
        Prepare row data for Excel export
        
        Args:
            file_info: Dictionary containing all file analysis information
            
        Returns:
            Dictionary formatted for Excel export
        """
        # Convert lists to strings for Excel compatibility
        bullet_points = file_info.get('bullet_points', [])
        if isinstance(bullet_points, list):
            bullet_points = '\n'.join([f"• {point}" for point in bullet_points])
        
        tags = file_info.get('tags', [])
        if isinstance(tags, list):
            tags = ', '.join(tags)
        
        return {
            'file_name': file_info.get('file_name', ''),
            'file_type': file_info.get('file_type', ''),
            'extracted_text': file_info.get('extracted_text', '')[:1000] + "..." if len(file_info.get('extracted_text', '')) > 1000 else file_info.get('extracted_text', ''),  # Truncate for Excel
            'summary': file_info.get('summary', ''),
            'description': file_info.get('description', ''),
            'bullet_points': bullet_points,
            'version': file_info.get('version', 'N/A'),
            'tags': tags,
            'duplicates': file_info.get('duplicate_reason', 'No duplicates found'),
            'duplicates_percentage': file_info.get('similarity_percentage', 0.0),
            'master_one': file_info.get('master_file', file_info.get('file_name', ''))
        }